from __future__ import annotations

import json
import os

from app import config
from app.models import TranscriptResult, Segment
from app.stitch import humanize_speaker, merge_speaker_runs


def _prose_segments(result: TranscriptResult) -> list[Segment]:
    """Сегменты для читаемого протокола (TXT/MD/DOCX): при диаризации склеиваем
    подряд идущие реплики одного спикера в одну (с порогом config.MERGE_MAX_SECONDS,
    чтобы длинный монолог резался на под-блоки); иначе — как есть (Whisper-сегменты).
    Субтитры (SRT/VTT) и JSON используют исходные мелкие сегменты."""
    if not result.diarized:
        return result.segments
    return merge_speaker_runs(result.segments, max_seconds=config.MERGE_MAX_SECONDS)


def format_timestamp(seconds: float, sep: str = ",") -> str:
    if seconds < 0:
        seconds = 0.0
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


def _hms(seconds: float) -> str:
    return format_timestamp(seconds, ",")[:8]


def _seg_label(seg: Segment) -> str:
    return humanize_speaker(seg.speaker)


def _prefix(seg: Segment, diarized: bool) -> str:
    """Возвращает 'Спикер N: ' если диаризация выполнена, иначе пустую строку."""
    return f"{_seg_label(seg)}: " if diarized else ""


def to_txt(result: TranscriptResult) -> str:
    lines = []
    for seg in _prose_segments(result):
        lines.append(f"[{_hms(seg.start)}] {_prefix(seg, result.diarized)}{seg.text.strip()}")
    return "\n".join(lines) + "\n"


def to_srt(result: TranscriptResult) -> str:
    blocks = []
    for i, seg in enumerate(result.segments, 1):
        ts = f"{format_timestamp(seg.start, ',')} --> {format_timestamp(seg.end, ',')}"
        text = f"{_prefix(seg, result.diarized)}{seg.text.strip()}"
        blocks.append(f"{i}\n{ts}\n{text}\n")
    return "\n".join(blocks)


def to_vtt(result: TranscriptResult) -> str:
    blocks = ["WEBVTT\n"]
    for seg in result.segments:
        ts = f"{format_timestamp(seg.start, '.')} --> {format_timestamp(seg.end, '.')}"
        text = f"{_prefix(seg, result.diarized)}{seg.text.strip()}"
        blocks.append(f"{ts}\n{text}\n")
    return "\n".join(blocks)


def to_json(result: TranscriptResult) -> str:
    return json.dumps(result.to_dict(), ensure_ascii=False, indent=2)


def to_md(result: TranscriptResult) -> str:
    lines = [f"# Транскрипция встречи\n",
             f"- Язык: {result.language}",
             f"- Модель: {result.model}",
             f"- Диаризация: {'да' if result.diarized else 'нет'}\n"]
    for seg in _prose_segments(result):
        pfx = _prefix(seg, result.diarized)
        header = f"**[{_hms(seg.start)}] {pfx.rstrip()}**" if pfx else f"**[{_hms(seg.start)}]**"
        lines.append(f"{header} {seg.text.strip()}\n")
    return "\n".join(lines)


def to_docx(result: TranscriptResult, path: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading("Транскрипция встречи", level=1)
    doc.add_paragraph(f"Язык: {result.language} · Модель: {result.model} · "
                      f"Диаризация: {'да' if result.diarized else 'нет'}")
    for seg in _prose_segments(result):
        p = doc.add_paragraph()
        bold_part = f"[{_hms(seg.start)}] {_prefix(seg, result.diarized)}"
        p.add_run(bold_part).bold = True
        p.add_run(seg.text.strip())
    doc.save(path)


def write_all(result: TranscriptResult, out_dir: str, formats: list[str],
              basename: str) -> list[str]:
    os.makedirs(out_dir, exist_ok=True)
    # JSON пишется ВСЕГДА, независимо от галочек форматов: это машинный источник
    # для стадии analyze (её вход — output/<job_id>/<basename>.json).
    formats = list(formats)
    if "json" not in formats:
        formats.append("json")
    written: list[str] = []
    text_map = {"txt": to_txt, "srt": to_srt, "vtt": to_vtt, "json": to_json, "md": to_md}
    for fmt in formats:
        path = os.path.join(out_dir, f"{basename}.{fmt}")
        if fmt in text_map:
            with open(path, "w", encoding="utf-8") as f:
                f.write(text_map[fmt](result))
            written.append(path)
        elif fmt == "docx":
            to_docx(result, path)
            written.append(path)
    return written
